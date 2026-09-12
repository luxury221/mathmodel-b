import importlib
import importlib.metadata
import json
import logging
import os
import platform
import sys
import tempfile
import threading
import time
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

LOGGER = logging.getLogger(__name__)
ROOT = Path(__file__).resolve().parents[1]
REPORT_DIRECTORY = ROOT / 'reports' / 'environment_smoke'
PATH_VARIABLES = (
    'TEMP', 'TMP', 'PIP_CACHE_DIR', 'UV_CACHE_DIR', 'UV_PYTHON_INSTALL_DIR',
    'JUPYTER_CONFIG_DIR', 'JUPYTER_DATA_DIR', 'JUPYTER_RUNTIME_DIR',
    'IPYTHONDIR', 'MPLCONFIGDIR', 'NUMBA_CACHE_DIR', 'JOBLIB_TEMP_FOLDER',
    'XDG_CACHE_HOME', 'XDG_DATA_HOME', 'PYTHONPYCACHEPREFIX',
)
IMPORT_NAMES = {
    'scikit-learn': 'sklearn',
    'pillow': 'PIL',
    'pyyaml': 'yaml',
    'python-docx': 'docx',
    'pymupdf': 'pymupdf',
}


def require(condition, message):
    if not condition:
        raise RuntimeError(message)


def check_storage():
    require(ROOT.drive.upper() == 'D:', 'Workspace is not on drive D.')
    require(Path(sys.executable).is_relative_to(ROOT / '.venv'), 'Wrong Python interpreter.')
    require(Path(sys.base_prefix).is_relative_to(ROOT / '.runtime'), 'Python runtime is not project-local.')
    expected_version = (ROOT / '.python-version').read_text(encoding='utf-8').strip()
    require(platform.python_version() == expected_version, 'Python version differs from the pinned version.')
    paths = {name: os.environ.get(name, '') for name in PATH_VARIABLES}
    for name, value in paths.items():
        require(value and Path(value).drive.upper() == 'D:', f'{name} must point to drive D.')
    require(Path(tempfile.gettempdir()).drive.upper() == 'D:', 'Temporary files would be written outside D.')
    return {'python': sys.executable, 'base_prefix': sys.base_prefix, 'paths': paths}


def check_imports():
    versions = {}
    for distribution in (ROOT / 'requirements.in').read_text(encoding='utf-8').splitlines():
        distribution = distribution.strip()
        if not distribution:
            continue
        importlib.import_module(IMPORT_NAMES.get(distribution, distribution))
        versions[distribution] = importlib.metadata.version(distribution)
    return versions


def check_numerical_libraries():
    import networkx
    import numpy
    import scipy.linalg
    import scipy.optimize
    import sympy
    from numba import njit

    matrix = numpy.array([[3.0, 1.0], [1.0, 2.0]])
    right_hand_side = numpy.array([9.0, 8.0])
    solution = scipy.linalg.solve(matrix, right_hand_side)
    require(numpy.allclose(solution, [2.0, 3.0]), 'Linear algebra check failed.')
    result = scipy.optimize.minimize(lambda values: ((values - 2.0) ** 2).sum(), numpy.zeros(2))
    require(result.success and numpy.allclose(result.x, 2.0, atol=1e-5), 'SciPy optimization check failed.')
    symbol = sympy.Symbol('sample')
    require(sympy.diff(symbol ** 3, symbol) == 3 * symbol ** 2, 'Symbolic calculation failed.')
    graph = networkx.path_graph(4)
    require(networkx.shortest_path_length(graph, 0, 3) == 3, 'Graph library check failed.')
    compiled_sum = njit(cache=True)(sum_sample_values)
    require(compiled_sum(numpy.arange(5.0)) == 10.0, 'Numba native compilation failed.')
    return {'linear_solution': solution.tolist(), 'numba_native_execution': True}


def sum_sample_values(values):
    return values.sum()


def check_geometry():
    import shapely
    from shapely.geometry import box

    intersection = box(0, 0, 2, 2).intersection(box(1, 1, 3, 3))
    require(abs(intersection.area - 1.0) < 1e-10, 'GEOS geometry operation failed.')
    return {'geos_version': shapely.geos_version_string, 'intersection_area': intersection.area}


def check_solvers():
    import cvxpy
    from ortools.sat.python import cp_model

    decision = cvxpy.Variable()
    problem = cvxpy.Problem(cvxpy.Minimize(cvxpy.square(decision - 2)), [decision >= 0])
    problem.solve(solver='CLARABEL')
    require(problem.status == 'optimal' and abs(float(decision.value) - 2.0) < 1e-4, 'CVXPY solver failed.')
    model = cp_model.CpModel()
    integer_decision = model.new_int_var(0, 9, 'environment_check')
    model.maximize(integer_decision)
    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = 5
    status = solver.solve(model)
    require(status == cp_model.OPTIMAL and solver.value(integer_decision) == 9, 'OR-Tools native solver failed.')
    return {'cvxpy_solvers': cvxpy.installed_solvers(), 'ortools_cp_sat': 'passed'}


def check_plotting():
    import matplotlib

    matplotlib.use('Agg')
    from matplotlib import font_manager, pyplot

    font_path = font_manager.findfont(font_manager.FontProperties(family='Microsoft YaHei'), fallback_to_default=False)
    pyplot.rcParams['font.sans-serif'] = ['Microsoft YaHei']
    pyplot.rcParams['axes.unicode_minus'] = False
    figure, axes = pyplot.subplots(figsize=(6, 3.5))
    axes.plot([0, 1, 2], [0, 1, 0], marker='o')
    axes.set_title('环境自检（非赛题实验结果）')
    axes.set_xlabel('测试序号')
    axes.set_ylabel('示例数值')
    figure.tight_layout()
    image_path = REPORT_DIRECTORY / 'plot_check.png'
    pdf_path = REPORT_DIRECTORY / 'plot_check.pdf'
    figure.savefig(image_path, dpi=140)
    figure.savefig(pdf_path)
    pyplot.close(figure)
    require(image_path.stat().st_size > 1000 and pdf_path.stat().st_size > 1000, 'Figure export failed.')
    return {'chinese_font': font_path, 'png': str(image_path), 'pdf': str(pdf_path)}


def check_document_io():
    import pandas
    from docx import Document
    from pypdf import PdfReader
    from reportlab.pdfgen import canvas

    frame = pandas.DataFrame({'sample_id': [1, 2], 'value': [3.0, 4.0]})
    workbook_path = REPORT_DIRECTORY / 'spreadsheet_check.xlsx'
    frame.to_excel(workbook_path, index=False)
    restored = pandas.read_excel(workbook_path)
    require(restored.shape == (2, 2), 'Excel round-trip failed.')
    document_path = REPORT_DIRECTORY / 'document_check.docx'
    document = Document()
    document.add_paragraph('Environment diagnostic only; not a competition result.')
    document.save(document_path)
    require(len(Document(document_path).paragraphs) == 1, 'Word round-trip failed.')
    pdf_path = REPORT_DIRECTORY / 'document_check.pdf'
    pdf_canvas = canvas.Canvas(str(pdf_path))
    pdf_canvas.drawString(60, 760, 'Environment diagnostic only; not a competition result.')
    pdf_canvas.save()
    require('Environment diagnostic' in PdfReader(pdf_path).pages[0].extract_text(), 'PDF round-trip failed.')
    source_pages = len(PdfReader(ROOT / 'B题.pdf').pages)
    source_paragraphs = {path.name: len(Document(path).paragraphs) for path in (ROOT / '附件').glob('*.docx')}
    return {'source_pdf_pages': source_pages, 'source_word_paragraphs': source_paragraphs}


class DiagnosticHandler(BaseHTTPRequestHandler):
    def do_GET(self):
        payload = json.dumps({'status': 'ok', 'purpose': 'environment_diagnostic'}).encode('utf-8')
        self.send_response(200)
        self.send_header('Content-Type', 'application/json')
        self.send_header('Content-Length', str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)

    def log_message(self, message_format, *arguments):
        return


def check_local_http():
    import requests

    server = ThreadingHTTPServer(('127.0.0.1', 0), DiagnosticHandler)
    worker = threading.Thread(target=server.serve_forever, daemon=True)
    worker.start()
    try:
        with requests.Session() as session:
            session.trust_env = False
            response = session.get(f'http://127.0.0.1:{server.server_port}/health', timeout=5)
        require(response.status_code == 200 and response.json()['status'] == 'ok', 'Local HTTP/JSON check failed.')
        return {'bind_address': '127.0.0.1', 'official_simulator_port_contacted': False}
    finally:
        server.shutdown()
        server.server_close()
        worker.join(timeout=5)


def check_jupyter_kernel():
    from jupyter_client import KernelManager
    from jupyter_client.kernelspec import KernelSpecManager

    specification = KernelSpecManager().get_kernel_spec('b2026')
    require(Path(specification.argv[0]).is_relative_to(ROOT / '.venv'), 'Jupyter selected the wrong interpreter.')
    require(specification.env.get('TEMP', '').upper().startswith('D:'), 'Kernel does not preserve D-drive temporary files.')
    manager = KernelManager(kernel_name='b2026')
    client = None
    output = []
    try:
        manager.start_kernel(cwd=str(ROOT))
        client = manager.client()
        client.start_channels()
        client.wait_for_ready(timeout=45)
        request_id = client.execute("import json, sys, tempfile; print(json.dumps({'python':sys.executable,'temp':tempfile.gettempdir()}))")
        deadline = time.monotonic() + 45
        while time.monotonic() < deadline:
            message = client.get_iopub_msg(timeout=10)
            if message.get('parent_header', {}).get('msg_id') != request_id:
                continue
            if message['msg_type'] == 'stream':
                output.append(message['content']['text'])
            if message['msg_type'] == 'error':
                raise RuntimeError(message['content']['evalue'])
            if message['msg_type'] == 'status' and message['content']['execution_state'] == 'idle':
                break
        result = json.loads(''.join(output).strip())
        require(Path(result['python']).is_relative_to(ROOT / '.venv'), 'Kernel executed outside the isolated environment.')
        require(Path(result['temp']).drive.upper() == 'D:', 'Kernel temporary directory is not on D.')
        return result
    finally:
        if client is not None:
            client.stop_channels()
        if manager.has_kernel:
            manager.shutdown_kernel(now=True)


def simulator_status():
    manifest_path = ROOT / 'simulator' / 'installation.json'
    if not manifest_path.exists():
        return {'status': 'simulator_not_installed', 'reason': 'Local installation manifest is missing.', 'official_interface_actions_sent': 0}
    manifest = json.loads(manifest_path.read_text(encoding='utf-8-sig'))
    executable = ROOT / 'simulator' / manifest['executable']
    if not executable.is_file():
        status = 'executable_missing'
    elif manifest.get('startup_verified'):
        status = 'startup_verified_login_not_checked'
    else:
        status = 'files_present_startup_not_verified'
    return {
        'status': status,
        'executable': str(executable),
        'startup_verified_at': manifest.get('startup_verified_at'),
        'startup_report': manifest.get('startup_report'),
        'login_verified_by_this_check': False,
        'official_interface_actions_sent': 0,
    }


def main():
    report = {'checked_at': datetime.now().astimezone().isoformat(), 'checks': {}, 'failures': [], 'simulator': simulator_status()}
    try:
        report['checks']['storage'] = check_storage()
    except Exception as error:
        LOGGER.exception('Storage validation failed')
        print(f'FAIL storage: {error}')
        return 1
    REPORT_DIRECTORY.mkdir(parents=True, exist_ok=True)
    checks = (
        ('imports', check_imports),
        ('numerical_libraries', check_numerical_libraries),
        ('geometry', check_geometry),
        ('solvers', check_solvers),
        ('plotting', check_plotting),
        ('document_io', check_document_io),
        ('local_http', check_local_http),
        ('jupyter_kernel', check_jupyter_kernel),
    )
    for name, function in checks:
        try:
            report['checks'][name] = function()
            print(f'PASS {name}')
        except Exception as error:
            LOGGER.exception('Environment check failed: %s', name)
            report['failures'].append({'check': name, 'error': f'{type(error).__name__}: {error}'})
            print(f'FAIL {name}: {error}')
    if report['failures']:
        report['overall_status'] = 'environment_checks_failed'
    elif report['simulator']['status'] == 'startup_verified_login_not_checked':
        report['overall_status'] = 'local_environment_ready_login_not_checked'
    else:
        report['overall_status'] = 'python_environment_ready'
    report_path = ROOT / 'reports' / 'environment.json'
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f"Simulator status: {report['simulator']['status']}")
    print(f'Report: {report_path}')
    return 1 if report['failures'] else 0


if __name__ == '__main__':
    raise SystemExit(main())
